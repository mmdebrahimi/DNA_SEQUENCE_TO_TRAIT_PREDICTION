"""Load the Agarwal 2025 human lentiMPRA panel: 230 bp sequences with measured log2(RNA/DNA).

SUBSTRATE, measured 2026-10-27 from the verified deposit (Zenodo 10558183, CC-BY-4.0):

| cell  | designed (S3) | measured (S6) | JOINED        | usable |
|-------|---------------|---------------|---------------|--------|
| HepG2 | 164,307       | 61,463        | 61,463 (1.000)| yes    |
| K562  | 243,778       | 98,332        | 98,332 (1.000)| yes    |
| WTC11 |  75,542       | 23,064        |      0 (0.000)| **NO** |

So the usable labelled set is **159,795 pairs** over two cell types, every sequence exactly 230 bp,
with the authors' own **10-fold assignment** carried through.

WHY WTC11 REFUSES RATHER THAN RETURNING WHAT IT CAN. Its S6 activity rows are keyed by ENSEMBL GENE
ID (`ENSG00000000003`) while its S3 design rows use element names (`seq1000_F`) -- different entity
types, so the intersection is empty. Table S7 (the joint library) rescues only **699 of 23,064
(3.03%)**, because S7's ENSG-style names are its 2,396 `promoter` rows and nothing else. Returning
699 rows would hand a caller a 3%-complete cohort that looks like a cohort; this repo has paid for
that shape before (the sigma70 rung's naive key join returned 3,821 of 10,898, which is why
`expression.load_pairs` enforces MIN_JOIN_RATE). So the join rate is ENFORCED here too.

TWO PARSING TRAPS, both measured rather than assumed:

1. **The header row MOVES between sheets of the same workbook** -- S3 has it at row 4 (HepG2),
   row 2 (K562) and row 1 (WTC11), because some sheets carry a free-text design note above it. A
   fixed `min_row` reads a 327-character description as a column name and then finds no columns. So
   the header is DISCOVERED by looking for the tokens, and the row it was found at is REPORTED.
2. **The 230 nt includes the cloning adapters** -- the column is named
   `230nt sequence (15nt 5' adaptor - ...)`. The biological insert is shorter than 230. This module
   returns the sequence AS PUBLISHED and records the verbatim column name in `LentiStats`, so a
   caller trimming adapters does so deliberately. It does not trim silently, because a silent trim
   would change what every downstream number is about.

PUBLISHED INCUMBENTS live in the same workbook (`Benchmarking results`): per-fold figures for
Biochemical / MPRAnn / Sei / Enformer / MPRALegNet. **The sheet does not state the METRIC**, so
`published_incumbents()` returns the numbers with `metric="UNSPECIFIED_IN_SOURCE"` and callers must
not compare a Spearman of their own against them until that is resolved -- a readout-type mismatch
is not a comparison.
"""
from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path

#: The deposit lands here, namespace-separate from the E. coli `mpra/` directory so a human table
#: can never be mistaken for a GSE144621 one (the shared-key silent-overwrite trap).
DEFAULT_DIR = Path("D:/dna_decode_cache/mpra_human")

S3_FILE = "Table S3 - large-scale lib design.xlsx"
S6_FILE = "Table S6 - folds and performance.xlsx"

#: Mirrors `expression.load_pairs`. A cohort below this is a key problem, not a small cohort.
MIN_JOIN_RATE = 0.95

#: Measured: all 483,627 designed sequences across all three sheets are exactly this long.
EXPECTED_LEN = 230

CELL_TYPES = ("HepG2", "K562", "WTC11")


class LentiMpraDataError(ValueError):
    """Refuse rather than return a plausible-looking partial cohort."""


#: Parsed-workbook cache, keyed by (kind, cell, data_dir). Table S3 is 56 MB and openpyxl takes
#: tens of seconds per read_only pass over it, so an uncached loader makes any caller that touches
#: two cell types -- or a test module that touches several -- pay minutes. Measured: the first
#: version of `tests/test_lentimpra.py` exceeded a 600 s timeout purely by re-parsing.
#:
#: Values are stored as TUPLES so a cached payload cannot be mutated by one caller and observed
#: changed by the next; each public call rebuilds its own dict from them.
_CACHE: dict = {}


def clear_cache() -> None:
    """Drop the parsed-workbook cache. For tests that swap the data directory under the loader."""
    _CACHE.clear()


@dataclass(frozen=True)
class LentiRecord:
    name: str
    seq: str
    activity: float          # observed log2(RNA/DNA)
    fold: int                # the authors' published fold, 1..10
    cell: str


@dataclass
class LentiStats:
    cell: str = ""
    n_designed: int = 0
    n_measured: int = 0
    n_joined: int = 0
    join_rate: float = 0.0
    s3_sheet: str = ""
    s3_header_row: int = 0
    s6_sheet: str = ""
    s6_header_row: int = 0
    sequence_column: str = ""     # verbatim, because it records that adapters are included
    length_min: int = 0
    length_max: int = 0
    n_folds: int = 0
    fold_sizes: dict = field(default_factory=dict)

    def as_dict(self) -> dict:
        d = {k: getattr(self, k) for k in (
            "cell", "n_designed", "n_measured", "n_joined", "s3_sheet", "s3_header_row",
            "s6_sheet", "s6_header_row", "sequence_column", "length_min", "length_max", "n_folds")}
        d["join_rate"] = round(self.join_rate, 4)
        d["fold_sizes"] = dict(sorted(self.fold_sizes.items()))
        return d


def _find_header(ws, tokens, limit: int = 8):
    """First row (1-based) containing every token, plus its cells. The row MOVES; see module doc."""
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=limit, values_only=True), start=1):
        cells = ["" if c is None else str(c).strip() for c in row]
        low = [c.lower() for c in cells]
        if all(any(tok in c for c in low) for tok in tokens):
            return i, cells
    raise LentiMpraDataError(
        "no header row containing %s in the first %d rows; the sheet layout changed and a fixed "
        "row index would read a description string as a column name" % (list(tokens), limit))


def _sheet_for(wb, cell: str, *, suffix: bool):
    want = cell.lower()
    for s in wb.sheetnames:
        sl = s.lower()
        if (sl.endswith(want) if suffix else sl.startswith(want)):
            return s
    raise LentiMpraDataError("no sheet for cell %r among %s" % (cell, wb.sheetnames))


def load_design(cell: str, data_dir: Path | str = DEFAULT_DIR) -> tuple[dict, LentiStats]:
    """`{name: seq}` from Table S3, for one cell type."""
    import openpyxl

    st = LentiStats(cell=cell)
    path = Path(data_dir) / S3_FILE
    if not path.is_file():
        raise LentiMpraDataError(
            "%s is absent. Fetch it with `uv run python scripts/fetch_lentimpra.py "
            "--select tables`, which verifies every byte against the publisher's MD5." % path)

    key = ("design", cell, str(path))
    if key in _CACHE:
        pairs, sheet, hrow, seq_col, lmin, lmax = _CACHE[key]
        st.s3_sheet, st.s3_header_row, st.sequence_column = sheet, hrow, seq_col
        st.n_designed, st.length_min, st.length_max = len(pairs), lmin, lmax
        return dict(pairs), st

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        st.s3_sheet = _sheet_for(wb, cell, suffix=False)
        ws = wb[st.s3_sheet]
        st.s3_header_row, hdr = _find_header(ws, ("name", "sequence"))
        j_name = next(j for j, h in enumerate(hdr) if h.lower() == "name")
        j_seq = next(j for j, h in enumerate(hdr) if "sequence" in h.lower())
        st.sequence_column = hdr[j_seq]
        out, lens = {}, []
        for row in ws.iter_rows(min_row=st.s3_header_row + 1, values_only=True):
            if j_name >= len(row) or j_seq >= len(row):
                continue
            nm, sq = row[j_name], row[j_seq]
            if nm is None or sq is None:
                continue
            nm, sq = str(nm).strip(), str(sq).strip().upper()
            if not nm or not re.fullmatch(r"[ACGTN]+", sq):
                continue
            out[nm] = sq
            lens.append(len(sq))
        if not out:
            raise LentiMpraDataError(
                "%s/%s parsed ZERO sequences. A zero-row parse that is then joined reads as 'no "
                "overlap' rather than 'no data', so this refuses instead." % (path.name, st.s3_sheet))
        st.n_designed = len(out)
        st.length_min, st.length_max = min(lens), max(lens)
        _CACHE[key] = (tuple(out.items()), st.s3_sheet, st.s3_header_row,
                       st.sequence_column, st.length_min, st.length_max)
        return out, st
    finally:
        wb.close()


def load_activity(cell: str, data_dir: Path | str = DEFAULT_DIR) -> tuple[dict, LentiStats]:
    """`{sequence_id: (activity, fold)}` from Table S6, for one cell type."""
    import openpyxl

    st = LentiStats(cell=cell)
    path = Path(data_dir) / S6_FILE
    if not path.is_file():
        raise LentiMpraDataError("%s is absent; see scripts/fetch_lentimpra.py" % path)

    key = ("activity", cell, str(path))
    if key in _CACHE:
        pairs, sheet, hrow = _CACHE[key]
        st.s6_sheet, st.s6_header_row, st.n_measured = sheet, hrow, len(pairs)
        return dict(pairs), st

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        st.s6_sheet = _sheet_for(wb, cell, suffix=True)
        ws = wb[st.s6_sheet]
        st.s6_header_row, hdr = _find_header(ws, ("sequence id", "observed"))
        low = [h.lower() for h in hdr]
        j_id = next(j for j, h in enumerate(low) if "sequence id" in h)
        j_y = next(j for j, h in enumerate(low) if "observed" in h)
        j_f = next((j for j, h in enumerate(low) if h == "fold"), None)
        if j_f is None:
            raise LentiMpraDataError(
                "%s has no `Fold` column; the published fold assignment is what makes this cohort "
                "leakage-safe and comparable to the published benchmark, so it is required"
                % st.s6_sheet)
        out = {}
        for row in ws.iter_rows(min_row=st.s6_header_row + 1, values_only=True):
            if max(j_id, j_y, j_f) >= len(row):
                continue
            sid, y, f = row[j_id], row[j_y], row[j_f]
            if sid is None or y is None or f is None:
                continue
            try:
                out[str(sid).strip()] = (float(y), int(f))
            except (TypeError, ValueError):
                continue
        if not out:
            raise LentiMpraDataError("%s parsed ZERO activities" % st.s6_sheet)
        st.n_measured = len(out)
        _CACHE[key] = (tuple(out.items()), st.s6_sheet, st.s6_header_row)
        return out, st
    finally:
        wb.close()


def load_cell(cell: str, data_dir: Path | str = DEFAULT_DIR, *,
              min_join_rate: float = MIN_JOIN_RATE) -> tuple[list, LentiStats]:
    """Join design to activity for one cell type. REFUSES below `min_join_rate`.

    The join rate is computed over the MEASURED side: the design library is deliberately larger than
    the measured set (not every designed element yields a usable readout), so measured-coverage is
    the meaningful denominator and design-coverage is reported but never gated on.
    """
    if cell not in CELL_TYPES:
        raise LentiMpraDataError("unknown cell %r; have %s" % (cell, list(CELL_TYPES)))
    seqs, st3 = load_design(cell, data_dir)
    acts, st6 = load_activity(cell, data_dir)

    st = LentiStats(
        cell=cell, n_designed=st3.n_designed, n_measured=st6.n_measured,
        s3_sheet=st3.s3_sheet, s3_header_row=st3.s3_header_row,
        s6_sheet=st6.s6_sheet, s6_header_row=st6.s6_header_row,
        sequence_column=st3.sequence_column,
        length_min=st3.length_min, length_max=st3.length_max,
    )

    records, lens = [], []
    for sid, (y, fold) in acts.items():
        sq = seqs.get(sid)
        if sq is None:
            continue
        records.append(LentiRecord(name=sid, seq=sq, activity=y, fold=fold, cell=cell))
        lens.append(len(sq))

    st.n_joined = len(records)
    st.join_rate = len(records) / st.n_measured if st.n_measured else 0.0
    for r in records:
        st.fold_sizes[r.fold] = st.fold_sizes.get(r.fold, 0) + 1
    st.n_folds = len(st.fold_sizes)
    if lens:
        st.length_min, st.length_max = min(lens), max(lens)

    if st.join_rate < min_join_rate:
        raise LentiMpraDataError(
            "%s joins at %.4f (%d of %d measured rows), below the %.2f bar. This is a KEY problem, "
            "not a small cohort: S6 keys this cell by Ensembl gene id while S3 uses element names. "
            "Returning the overlap would hand back a %.1f%%-complete cohort that looks complete."
            % (cell, st.join_rate, st.n_joined, st.n_measured, min_join_rate, 100 * st.join_rate))

    if st.length_min != st.length_max:
        raise LentiMpraDataError(
            "ragged sequences (%d-%d bp) where the panel is documented uniform; a downstream "
            "fixed-width model would silently pad or truncate" % (st.length_min, st.length_max))
    if st.length_max != EXPECTED_LEN:
        raise LentiMpraDataError(
            "sequences are %d bp, expected %d. The deposit may have changed; refusing rather than "
            "proceeding on a different substrate." % (st.length_max, EXPECTED_LEN))
    return records, st


def usable_cells(data_dir: Path | str = DEFAULT_DIR) -> dict:
    """Which cell types clear the join bar. Reports the failures rather than hiding them."""
    out = {}
    for cell in CELL_TYPES:
        try:
            recs, st = load_cell(cell, data_dir)
            out[cell] = {"usable": True, "n": len(recs), "stats": st.as_dict()}
        except LentiMpraDataError as exc:
            out[cell] = {"usable": False, "reason": str(exc)[:220]}
    return out


def published_incumbents(data_dir: Path | str = DEFAULT_DIR) -> dict:
    """Per-fold published figures for the five benchmarked methods, from Table S6.

    `metric` is **UNSPECIFIED_IN_SOURCE**: the sheet gives numbers and method names and never says
    whether they are Pearson, Spearman or something else. Comparing a Spearman of one's own against
    an unknown readout is a readout-type mismatch, not a comparison, so the field says so instead of
    guessing the likely answer.
    """
    import openpyxl

    path = Path(data_dir) / S6_FILE
    if not path.is_file():
        raise LentiMpraDataError("%s is absent; see scripts/fetch_lentimpra.py" % path)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb["Benchmarking results"]
        blocks, cur, methods = {}, None, None
        for row in ws.iter_rows(values_only=True):
            cells = [c for c in row]
            first = "" if cells[0] is None else str(cells[0]).strip()
            if first in CELL_TYPES:
                cur = first
                blocks[cur] = {}
                methods = None
                continue
            if first.lower() == "test fold":
                methods = ["" if c is None else str(c).strip() for c in cells[1:]]
                continue
            if cur and methods and isinstance(cells[0], (int, float)):
                fold = int(cells[0])
                for m, v in zip(methods, cells[1:]):
                    if m and isinstance(v, (int, float)):
                        blocks[cur].setdefault(m, {})[fold] = float(v)
        return {"metric": "UNSPECIFIED_IN_SOURCE", "source": S6_FILE,
                "sheet": "Benchmarking results", "by_cell": blocks}
    finally:
        wb.close()
